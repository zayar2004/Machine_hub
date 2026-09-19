"""Excel import helpers — flexible column detection."""
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


ALLOWED_STATUSES = ("ACTIVE", "INACTIVE", "ARCHIVED")


@dataclass
class ParsedRow:
    row_num: int
    data: dict[str, Any]
    errors: list[str] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return not self.errors


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    header_errors: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return len(self.rows)

    @property
    def valid(self) -> list[ParsedRow]:
        return [r for r in self.rows if r.is_valid]

    @property
    def invalid(self) -> list[ParsedRow]:
        return [r for r in self.rows if not r.is_valid]


def _norm_header(value: Any) -> str:
    """Normalize header: strip, lowercase, spaces/dashes to underscore."""
    if value is None:
        return ""
    s = str(value).strip().lower()
    s = s.replace(" ", "_").replace("-", "_")
    s = "".join(c for c in s if c.isalnum() or c == "_")
    return s


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _norm_status(value: Any) -> str:
    s = _cell_str(value).upper()
    if not s:
        return "ACTIVE"
    return s


# ---------- Column alias mapping ----------
# Maps user-friendly column names → system column names
MACHINE_ALIASES = {
    # Code
    "machine_code": "machine_code",
    "machinecode": "machine_code",
    "code": "machine_code",
    "id": "machine_code",
    # Name
    "machine_name": "machine_name",
    "machinename": "machine_name",
    "name": "machine_name",
    "title": "machine_name",
    # Description
    "description": "description",
    "desc": "description",
    "location": "description",   # Accept "Location" as description
    "note": "description",
    "notes": "description",
    # Status
    "status": "status",
    # Ignored columns
    "no": "_ignore",
    "arrival_date": "_ignore",
    "date": "_ignore",
    "arrivaldate": "_ignore",
}


def parse_machines_excel(file_bytes: bytes) -> ParseResult:
    """Parse machine Excel with flexible column detection.

    Accepts:
      - machine_code (required, or 'Machine Code', 'Code', etc.)
      - machine_name (optional if machine_code exists — uses code as name)
      - description / location (optional)
      - status (optional)

    Ignores:
      - No, Arrival Date, and other unknown columns
    """
    result = ParseResult()

    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        result.header_errors.append(f"Could not open Excel file: {e}")
        return result

    ws = wb.active
    if ws is None:
        result.header_errors.append("Excel file has no sheets.")
        return result

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        result.header_errors.append("Excel file is empty.")
        return result

    # Normalize + map headers
    raw_headers = [_norm_header(h) for h in header_row]
    col_map: dict[str, int] = {}  # system_col_name -> column_index
    unknown_cols: list[str] = []

    for idx, raw in enumerate(raw_headers):
        if not raw:
            continue
        mapped = MACHINE_ALIASES.get(raw)
        if mapped == "_ignore":
            continue
        if mapped is None:
            # Try prefix match (e.g. "machine_name_x" → "machine_name")
            for alias, system in MACHINE_ALIASES.items():
                if system != "_ignore" and raw.startswith(alias):
                    mapped = system
                    break
        if mapped is None:
            unknown_cols.append(raw)
            continue
        # First occurrence wins (don't overwrite)
        if mapped not in col_map:
            col_map[mapped] = idx

    # Validate required columns
    if "machine_code" not in col_map:
        result.header_errors.append(
            "Missing required column: machine_code (or 'Machine Code', 'Code')"
        )
        return result

    # machine_name optional — if missing, will use machine_code
    has_name = "machine_name" in col_map

    seen_codes: set[str] = set()

    for i, row in enumerate(rows_iter, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue  # skip blank rows

        def get(col_name: str) -> Any:
            idx = col_map.get(col_name)
            if idx is None or idx >= len(row):
                return ""
            return row[idx]

        code = _cell_str(get("machine_code")).upper()
        name = _cell_str(get("machine_name")) if has_name else ""
        desc = _cell_str(get("description"))
        status = _norm_status(get("status"))

        # If no name, use code as name
        if not name:
            name = code

        data = {
            "machine_code": code,
            "machine_name": name,
            "description": desc or None,
            "status": status,
        }

        errors: list[str] = []

        if not code:
            errors.append("machine_code is empty")
        elif len(code) > 60:
            errors.append("machine_code too long (max 60)")

        if not name:
            errors.append("machine_name is empty")
        elif len(name) > 160:
            errors.append("machine_name too long (max 160)")

        if status not in ALLOWED_STATUSES:
            errors.append(f"invalid status '{status}'")

        if code and code in seen_codes:
            errors.append(f"duplicate machine_code in file: {code}")
        elif code:
            seen_codes.add(code)

        result.rows.append(ParsedRow(row_num=i, data=data, errors=errors))

    if unknown_cols:
        result.header_errors.append(
            f"Note: ignored unknown column(s): {', '.join(unknown_cols)}"
        )

    return result


# ---------- ERRORS (global) ----------

ERROR_ALIASES = {
    "error_code": "error_code",
    "errorcode": "error_code",
    "code": "error_code",
    "error_name": "error_name",
    "errorname": "error_name",
    "name": "error_name",
    "error_fix": "error_fix",
    "errorfix": "error_fix",
    "fix": "error_fix",
    "solution": "error_fix",
    "explanation": "explanation",
    "desc": "explanation",
    "description": "explanation",
    "status": "status",
}


def parse_errors_excel(file_bytes: bytes) -> ParseResult:
    """Parse error Excel with flexible column detection."""
    result = ParseResult()

    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        result.header_errors.append(f"Could not open Excel file: {e}")
        return result

    ws = wb.active
    if ws is None:
        result.header_errors.append("Excel file has no sheets.")
        return result

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        result.header_errors.append("Excel file is empty.")
        return result

    raw_headers = [_norm_header(h) for h in header_row]
    col_map: dict[str, int] = {}

    for idx, raw in enumerate(raw_headers):
        if not raw:
            continue
        mapped = ERROR_ALIASES.get(raw)
        if mapped and mapped not in col_map:
            col_map[mapped] = idx

    if "error_code" not in col_map:
        result.header_errors.append("Missing required column: error_code")
        return result
    if "error_name" not in col_map:
        result.header_errors.append("Missing required column: error_name")
        return result

    seen_codes: set[str] = set()

    for i, row in enumerate(rows_iter, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue

        def get(col_name: str) -> Any:
            idx = col_map.get(col_name)
            if idx is None or idx >= len(row):
                return ""
            return row[idx]

        code = _cell_str(get("error_code")).upper()
        name = _cell_str(get("error_name"))
        fix = _cell_str(get("error_fix"))
        expl = _cell_str(get("explanation"))
        status = _norm_status(get("status"))

        data = {
            "error_code": code,
            "error_name": name,
            "error_fix": fix or None,
            "explanation": expl or None,
            "status": status,
        }

        errors: list[str] = []

        if not code:
            errors.append("error_code is empty")
        elif len(code) > 60:
            errors.append("error_code too long (max 60)")

        if not name:
            errors.append("error_name is empty")
        elif len(name) > 200:
            errors.append("error_name too long (max 200)")

        if status not in ALLOWED_STATUSES:
            errors.append(f"invalid status '{status}'")

        if code and code in seen_codes:
            errors.append(f"duplicate error_code in file: {code}")
        elif code:
            seen_codes.add(code)

        result.rows.append(ParsedRow(row_num=i, data=data, errors=errors))

    return result
